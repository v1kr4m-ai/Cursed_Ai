"""Extract plain text from a .docx. Prints the text to stdout, nothing else."""
import argparse
import sys

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--file', required=True)
    args = ap.parse_args()

    import docx
    d = docx.Document(args.file)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    sys.stdout.write('\n\n'.join(parts))


if __name__ == '__main__':
    main()
